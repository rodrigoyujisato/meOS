#!/usr/bin/env python3
"""Seed vault notes from ``07_Inbox_Agent/.seed/seed.xlsx``.

Currently implements the ``pessoas`` sheet -> ``05_People/``. Dry-run by default;
pass ``--apply`` to write. Idempotent: an existing note is never overwritten
(same guarantee as ``VaultManager.ensure_person_note``), so re-running only fills
the gaps. ``related_*`` columns become wikilinks in the note body — forward
references are fine, Obsidian resolves them once the target note exists.

Usage:
    uv run --with openpyxl python scripts/seed_from_xlsx.py            # dry-run
    uv run --with openpyxl python scripts/seed_from_xlsx.py --apply
    uv run --with openpyxl python scripts/seed_from_xlsx.py --xlsx /path/to/seed.xlsx
"""
from __future__ import annotations

import argparse
import sys
import unicodedata
from pathlib import Path

from jinja2 import Template
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from rysos.config import settings  # noqa: E402
from rysos.vault.manager import VaultManager  # noqa: E402
from rysos.vault.templates import PERSON_TEMPLATE  # noqa: E402

# The four fixed life-pillar areas (see 07_Inbox_Agent/.seed/_LEIA-ME.md).
AREA_STEMS = {
    "Família & Relacionamentos": "Familia_e_Relacionamentos",
    "Negócios & Governança": "Negocios_e_Governanca",
    "Patrimônio & Finanças": "Patrimonio_e_Financas",
    "Saúde & Vitalidade": "Saude_e_Vitalidade",
}
STEM_TITLES = {stem: title for title, stem in AREA_STEMS.items()}


def _split(cell: object) -> list[str]:
    return [part.strip() for part in str(cell or "").split(";") if part.strip()]


def _area_link(value: str) -> str:
    v = (value or "").strip()
    if not v:
        return ""
    if v in AREA_STEMS:
        title, stem = v, AREA_STEMS[v]
    elif v in STEM_TITLES:
        title, stem = STEM_TITLES[v], v
    else:
        return v  # unknown area — leave as plain text, don't invent a link
    return f"[[02_Areas/{stem}|{title}]]"


def render_person(row: dict[str, str], vault: VaultManager) -> str:
    """PERSON_TEMPLATE frontmatter/body + a '🧭 Contexto (seed)' section carrying the
    columns the base template has no slot for (area, related_*, free-text notes)."""
    name = row["name"].strip()
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    slug = vault._slugify(ascii_name).lower().replace(" ", "-")
    seed_tags = _split(row.get("tags"))

    note = Template(PERSON_TEMPLATE).render(
        slug=slug,
        name=name,
        organization=(row.get("organization") or "").strip(),
        role=(row.get("role") or "").strip(),
        email=(row.get("email") or "").strip(),
    )
    if seed_tags:
        extra = "".join(f"  - {t}\n" for t in seed_tags)
        note = note.replace("  - stakeholder\n", "  - stakeholder\n" + extra, 1)

    ctx: list[str] = ["", "---", "", "## 🧭 Contexto (seed)", ""]
    area = _area_link(row.get("area", ""))
    if area:
        ctx.append(f"- **Área:** {area}")
    projs = _split(row.get("related_projects"))
    if projs:
        ctx.append("- **Projetos:** " + ", ".join(f"[[01_Projects/{p}|{p}]]" for p in projs))
    people = _split(row.get("related_people"))
    if people:
        ctx.append("- **Pessoas:** " + ", ".join(f"[[05_People/{p}|{p}]]" for p in people))
    notes = (row.get("notes") or "").strip()
    if notes:
        ctx += ["", notes]
    ctx.append("")
    return note.rstrip() + "\n" + "\n".join(ctx)


def seed_people(xlsx: Path, vault: VaultManager, apply: bool) -> int:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if "pessoas" not in wb.sheetnames:
        print(f"! no 'pessoas' sheet in {xlsx}")
        return 1
    ws = wb["pessoas"]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(rows)]

    people_dir = vault.vault_path / "05_People"
    created = skipped = 0
    for raw in rows:
        row = {header[i]: raw[i] for i in range(len(header)) if i < len(raw)}
        name = str(row.get("name") or "").strip()
        if not name:
            continue
        path = people_dir / f"{vault._sanitize_filename(name)}.md"
        if path.exists():
            print(f"  skip   {name}  (note already exists)")
            skipped += 1
            continue
        if apply:
            vault.write_file_atomic(path, render_person(row, vault))
            print(f"  create {name}  -> {path.relative_to(vault.vault_path)}")
        else:
            print(f"  would create {name}  -> 05_People/{path.name}")
        created += 1

    verb = "created" if apply else "would create"
    print(f"\n{verb}: {created}   skipped (existing): {skipped}")
    if not apply and created:
        print("re-run with --apply to write.")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--apply", action="store_true", help="write notes (default: dry-run)")
    ap.add_argument("--xlsx", type=Path, default=None, help="override seed.xlsx path")
    args = ap.parse_args()

    vault = VaultManager()
    xlsx = args.xlsx or (vault.vault_path / "07_Inbox_Agent" / ".seed" / "seed.xlsx")
    if not xlsx.exists():
        print(f"! seed workbook not found: {xlsx}")
        return 1

    print(f"vault : {vault.vault_path}")
    print(f"seed  : {xlsx}")
    print(f"mode  : {'APPLY' if args.apply else 'dry-run'}\n")
    return seed_people(xlsx, vault, args.apply)


if __name__ == "__main__":
    raise SystemExit(main())
